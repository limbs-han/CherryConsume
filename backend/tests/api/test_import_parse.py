"""엑셀 가져오기의 파일 읽기. DB 없이 돈다. 작업 005 설계 5f, 설계 문서 8절. S11, E30, E35"""

import io
import zipfile
from datetime import date, datetime, time
from pathlib import Path

import openpyxl
import pytest

from cherry_api.imports import Unreadable, find_header, parse_rows, read_table, signature

HEAD = ["이용일자", "이용시간", "가맹점명", "이용금액(원)", "할부", "승인번호", "취소여부"]
BODY = [
    ["2026.09.10", "12:30", "스타벅스 강남점", "4,500", "일시불", "11112222", ""],
    ["2026-09-11", "", "이마트 성수점", "120,000", "3개월", "33334444", ""],
    ["2026/09/12", "08:05:09", "스타벅스 강남점", "-4,500", "일시불", "11112222", "취소"],
    ["합계", "", "", "120,000", "", "", ""],
]
# 날짜만 있는 행은 시각이 없다. 취소는 음수거나 취소 여부에 "취소"가 있다. 합계 줄은 건너뛴다
WANT = [
    (date(2026, 9, 10), time(12, 30), "스타벅스 강남점", 4500, 1, "11112222", False),
    (date(2026, 9, 11), None, "이마트 성수점", 120000, 3, "33334444", False),
    (date(2026, 9, 12), time(8, 5, 9), "스타벅스 강남점", 4500, 1, "11112222", True),
]


def cell(*parts: int) -> datetime:
    """엑셀 칸의 날짜는 시간대가 없는 값이다. 서버가 한국 시간으로 읽는다"""
    return datetime(*parts)  # noqa: DTZ001


def rows_of(data: bytes, name: str) -> list[tuple]:
    table = read_table(data, name)
    start, mapping = find_header(table)
    out = parse_rows(table, start, mapping)
    assert [r.error for r in out] == [None] * len(out)
    return [(r.day, r.at, r.merchant, r.amount, r.installment_months, r.approval_no, r.cancel) for r in out]


def csv_bytes(encoding: str) -> bytes:
    lines = ["카드 이용 내역", "조회 기간 2026.09.01 ~ 2026.09.30", ",".join(HEAD)]
    lines += [",".join(f'"{v}"' for v in row) for row in BODY]
    return "\n".join(lines).encode(encoding)


def test_csv_in_utf8_and_cp949_with_title_rows():
    # 머리 줄 위에 제목 두 줄이 있다. 위에서 20줄 안에서 열 이름이 가장 많이 맞는 줄을 머리 줄로 본다
    assert rows_of(csv_bytes("utf-8-sig"), "내역.csv") == WANT
    assert rows_of(csv_bytes("cp949"), "내역.csv") == WANT


def test_xlsx_with_real_dates_and_numbers():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["카드 이용 내역"])
    ws.append(HEAD)
    ws.append([cell(2026, 9, 10), time(12, 30), "스타벅스 강남점", 4500, "일시불", "11112222", None])
    ws.append([cell(2026, 9, 11), None, "이마트 성수점", 120000.0, "03", "33334444", None])
    ws.append([cell(2026, 9, 12, 8, 5, 9), None, "스타벅스 강남점", -4500, "일시불", "11112222", "취소"])
    buf = io.BytesIO()
    wb.save(buf)
    assert rows_of(buf.getvalue(), "내역.xlsx") == WANT


def test_html_table_saved_as_xls():
    # 카드사 웹의 xls는 html 표인 경우가 많다
    cells = "".join(f"<th>{h}</th>" for h in HEAD)
    body = "".join("<tr>" + "".join(f"<td>{v}</td>" for v in row) + "</tr>" for row in BODY)
    html = f"<html><body><table><tr>{cells}</tr>{body}</table></body></html>"
    assert rows_of(html.encode("utf-8"), "내역.xls") == WANT


def test_unknown_headers_need_mapping_and_user_mapping_works():
    # E30. 사전에 없는 열 이름이면 None이다. 사용자가 짝지은 열 이름으로 읽고, 짝은 머리 줄 모양으로 다시 쓴다
    data = "날,곳,값\n2026.09.10,스타벅스,4500\n".encode()
    table = read_table(data, "x.csv")
    assert find_header(table) == (None, None)
    start, mapping = find_header(table, {"row": 0, "columns": {"date": 0, "merchant": 1, "amount": 2}})
    [row] = parse_rows(table, start, mapping)
    assert (row.day, row.merchant, row.amount, row.cancel) == (date(2026, 9, 10), "스타벅스", 4500, False)
    assert signature(table[0]) == signature([" 날 ", "곳", "값"])


def test_bad_rows_and_files():
    data = "이용일자,가맹점명,이용금액\n2026.13.40,스타벅스,4500\n2026.09.10,스타벅스,사천오백\n".encode()
    table = read_table(data, "x.csv")
    start, mapping = find_header(table)
    assert [r.error for r in parse_rows(table, start, mapping)] == ["날짜를 읽지 못했어요", "금액을 읽지 못했어요"]
    # 진짜 옛 xls는 읽지 않는다. xlsx나 csv로 저장해 올리게 한다
    with pytest.raises(Unreadable):
        read_table(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "x.xls")
    with pytest.raises(Unreadable):
        read_table(("이용일자,가맹점명,이용금액\n" + "2026.09.10,a,1\n" * 3100).encode(), "x.csv")
    # 한 열을 두 칸에 고르거나 꼭 필요한 칸을 빼면 받지 않는다
    for bad in ({"date": 0, "merchant": 1, "amount": 1}, {"date": 0, "merchant": 1}):
        with pytest.raises(Unreadable):
            find_header(read_table("날,곳,값".encode(), "x.csv"), {"row": 0, "columns": bad})


def test_approval_installment_interest_free_and_region():
    # 자리만 채운 승인번호는 버린다. "무이자 3개월"은 무이자, "3개월"만 있으면 무이자인지 모른다. 해외 여부 Y는 해외다
    head = "이용일자,가맹점명,이용금액,할부,승인번호,해외여부"
    lines = [
        "2026.09.10,a,1000,무이자 3개월,-,N",
        "2026.09.10,b,1000,3개월,00000000,Y",
        "2026.09.10,c,1000,일시불,0012345,",
    ]
    table = read_table(chr(10).join([head, *lines]).encode(), "x.csv")
    start, mapping = find_header(table)
    got = [
        (r.installment_months, r.interest_free, r.approval_no, r.overseas) for r in parse_rows(table, start, mapping)
    ]
    assert got == [(3, True, None, False), (3, None, None, True), (1, False, "0012345", False)]


def test_xlsx_size_bombs():
    # 시트 머리의 크기를 믿지 않는다. 100만 번째 행에 칸 하나가 있어도 정해 둔 만큼만 읽는다
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEAD)
    ws.append(["2026.09.10", "", "스타벅스", "4500", "", "", ""])
    ws.cell(row=1_000_000, column=16_000, value="x")
    buf = io.BytesIO()
    wb.save(buf)
    table = read_table(buf.getvalue(), "x.xlsx")
    assert len(table) == 2 and max(len(r) for r in table) <= 60
    # 풀면 30MB가 넘는 압축 파일은 열지 않는다
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/sharedStrings.xml", "a" * 31_000_000)
    with pytest.raises(Unreadable):
        read_table(bomb.getvalue(), "x.xlsx")


def test_other_zip_methods_and_trailing_minus():
    # 재검토 높음 2번. 보통 쓰는 두 방식이 아닌 압축은 목차의 크기와 상관없이 읽지 않는다
    bz = io.BytesIO()
    with zipfile.ZipFile(bz, "w", zipfile.ZIP_BZIP2) as z:
        z.writestr("xl/sharedStrings.xml", "a")
    with pytest.raises(Unreadable):
        read_table(bz.getvalue(), "x.xlsx")
    # 낮음 11번. "4,500-"도 취소다
    table = read_table('이용일자,가맹점명,이용금액\n2026.09.10,스타벅스,"4,500-"\n'.encode(), "x.csv")
    start, mapping = find_header(table)
    [row] = parse_rows(table, start, mapping)
    assert (row.amount, row.cancel) == (4500, True)


FIXTURES = Path(__file__).resolve().parents[3] / "app" / "test" / "fixtures" / "imports"


def test_committed_fixtures_read_the_same():
    # 작업 006 계획 단계 1의 3. Dart 시험이 읽을 파일을 지금 읽기 코드가 위 시험들과 같게 읽는다
    for name in ("utf8.csv", "cp949.csv", "dates_and_numbers.xlsx"):
        assert rows_of((FIXTURES / name).read_bytes(), "내역." + name.rsplit(".", 1)[1]) == WANT
    table = read_table((FIXTURES / "far_cell.xlsx").read_bytes(), "x.xlsx")
    assert len(table) == 2 and max(len(r) for r in table) <= 60
    for name in ("unzip_bomb.xlsx", "bzip2.xlsx"):
        with pytest.raises(Unreadable):
            read_table((FIXTURES / name).read_bytes(), "x.xlsx")
